"""Stage-3 forecast layer on top of delivered V3: scenario ensemble -> engine runs -> immutable forecast tables.

python -X utf8 -m tools.microtom_forecast --help

Two engines, one growth authority each:
- `pipeline` (route `p1_pipeline`): members are `tools.tomato_fsp_pipeline` runs (prepare_run -> run, which commits
  checkpoints). Results are read ONLY from `committed/step_NNNN` and the frozen `native_inputs`; `advance_interval`
  is never called. Multi-instance, canopy placement included.
- `v3_direct` (routes `seed_route`, `stage2_fixture`): single-plant V3 runs (`vtc_development.prepare` /
  `vtc_growth_bridge.run` / `vtc_development.collect`), labelled `single_plant_no_canopy` in the manifest.

Nothing here changes biological state, applies a second growth or allocation update, multiplies light ratios
into traced photons, fits a coefficient, or labels a run anything but `assumption_conditioned`. The ensemble
is a declared scenario set over registry rows that can act on the input state (B3), not a calibrated probability.
"""
from __future__ import annotations

import argparse
import json
import os
import statistics
import uuid
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from tools import microtom_axes as ax
from tools.microtom_environment import load_light_input, make_climate_csv, resolve_ratio_pointer
from tools.microtom_forecast_io import (ASSUMPTION_COLUMNS, EVALUATION_STATUS, EVENT_COLUMNS, INTERVAL_COLUMNS, MAP_COLUMNS,
                                        MEMBER_COLUMNS, MTG_COLUMNS, NO_REUSE_MARK, RUN_STATUS, TOPOLOGY_SCOPE, TRAJECTORY_COLUMNS, UNRESOLVED_COLUMNS, VALUE_STATUS,
                                        code_digest, finalize, forecast_identifier, new_forecast_dir, now_iso, write_table)
from tools.microtom_parameters import (REGISTRY_PATH, SCENARIO_LABELS, apply_overrides, build_request_overrides, eligibility, load_registry,
                                       registry_hash, resolve_engine_capabilities, retained_from_request, sample_member, scenario_set)
from tools.vtc_growth_io import ROOT, ContractError, aware, digest, read_json, require, rows, sha, write_csv, write_json

SCENARIO_SCHEMA = 'microtom-forecast-scenario-2'
ENGINES = ('pipeline', 'v3_direct')
ROUTES = ('seed_route', 'stage2_fixture', 'p1_pipeline')
ROUTE_ENGINE = {'seed_route': 'v3_direct', 'stage2_fixture': 'v3_direct', 'p1_pipeline': 'pipeline'}
ISSUE_CLASSES = ('quantitative', 'diagnostic')
QUANTITATIVE_RAYS = 4_800_000
SIGNIFICANT_DIGITS = 10  # numeric rounding of aggregated values; GroIMP summation order jitters at ~1e-17
CODE_FILES = [ROOT / 'tools' / n for n in ('microtom_forecast.py', 'microtom_forecast_io.py', 'microtom_forecast_eval.py', 'microtom_axes.py', 'fsp_axes.py', 'fsp_axes_segments.py',
                                           'microtom_parameters.py', 'microtom_environment.py', 'vtc_development.py',
                                           'vtc_development_io.py', 'vtc_development_native.rgg', 'vtc_development_observation.py',
                                           'vtc_growth_bridge.py', 'vtc_growth_io.py', 'tomato_fsp_pipeline.py', 'tomato_fsp_io.py',
                                           'tomato_fsp_canopy.py', 'tomato_fsp_acceptance.py')]
DEFAULT_LAMPS = [dict(lamp_id='L1', x_m=0, y_m=0, z_m=.6, power_umol_s=1200, source_kind='synthetic_assumed_point_LED')]
DEFAULT_LIGHT = dict(reference_ppfd_umol_m2_s=300, source_kind='user_assumed', measured_at='')
DEFAULT_POINTER = 'data/records/env_light_20261010/plant_slot_light_ratio_20261010.csv'
REGISTRY_XLSX = ROOT / 'data/records/plant_registry/개체기록_MicroTom_20261010.xlsx'
SOWING_COLUMNS = ('파종일', 'sowing_datetime', 'sowing_date')
ORGAN_QUANTITIES = {  # class -> (quantity, state column, factor, unit)
    'Leaflet': [('leaflet_area_mm2', 'area_m2', 1e6, 'mm2'), ('leaflet_length_mm', 'length_m', 1e3, 'mm'), ('leaflet_width_mm', 'width_m', 1e3, 'mm')],
    'Internode': [('internode_length_mm', 'length_m', 1e3, 'mm'), ('internode_width_mm', 'width_m', 1e3, 'mm')],
    'CompoundLeaf': [('attachment_x_mm', 'x_m', 1e3, 'mm'), ('attachment_y_mm', 'y_m', 1e3, 'mm'), ('attachment_z_mm', 'z_m', 1e3, 'mm'),
                     ('petiole_base_azimuth_deg', 'azimuth_deg', 1, 'deg'), ('compound_elevation_deg', 'elevation_deg', 1, 'deg')],
    'Petiole': [('petiole_length_mm', 'length_m', 1e3, 'mm')],
    'RachisSegment': [('rachis_segment_length_mm', 'length_m', 1e3, 'mm'), ('rachis_bend_alpha_deg', 'rachis_bend_alpha_deg', 1, 'deg')],
    'TrussAxis': [('truss_axis_length_mm', 'length_m', 1e3, 'mm')],
}
PLANT_QUANTITIES = [('stem_length_mm', 'mm'), ('height_top_leaf_axis_mm', 'mm'), ('structural_mass_gDM', 'gDM'), ('root_mass_gDM', 'gDM'),
                    ('compounds', 'count'), ('trusses', 'count')]
EVENT_CLASSES = {'new_compound': 'CompoundLeaf', 'new_internode': 'Internode', 'new_truss': 'TrussAxis'}
# B4: the engine parameters whose values decide each event under V3's thermal program; equal across members -> rule_deterministic
THERMAL_RULE_KEYS = ('SP_STOP_RULE', 'CADENCE', 'FIRST_LEAVES', 'PHYLLOCHRON_DD', 'GERMINATION_DD', 'BASE_C', 'EMERGENCE_DAYS')
PAIR_CLASSES = ('Leaflet', 'RachisSegment')
AXIS_RULE_KEYS = ('APPLY_DETERMINATE_STOP',)  # per-axis flag that decides whether an axis initiates at all
AXIS_HIDDEN_DEFAULT = dict(stem_mass_gDM=0.02, stem_width_m=0.002,
                          source_ref='declared forecast assumption for the hidden stem of each declared axis (same values as the engine owner fixtures); not measured')
# G3 (F9, user decision 2026-10-13): the main fixture is re-issued only when one of these becomes true; the manifest names the trigger.
RERUN_TRIGGERS = {
    'initial_issue': 'no re-run condition claimed (first issue on this input, or a software re-issue)',
    'measured_light_landed': ('tools/reference_dli.py stops using an assumed reference and takes the calibrated field from '
                              'handoff/CODEX_CLI_PROMPT_P2_LAMP_CALIBRATION_2026-10-13.md / data/records/env_20261013/; gross assimilation then rests on '
                              'measurement and carbon limitation appears or does not AS A RESULT'),
    'real_p1_package': 'a real P1 package (real plant geometry) replaces the fixture; the engine has never run from real geometry so far',
    'nutrient_feedback_enabled': ('nutrient feedback (handoff/CODEX_CLI_PROMPT_P4_NUTRIENT_RESERVOIR_2026-10-12.md) is implemented and ENABLE_NUTRIENT_FEEDBACK '
                                  'is true, so f_nut multiplies gross assimilation and can itself create the limited state (out of scope on 2026-10-13)'),
}
HEADLINE = ('Conditional: after the main apex stops (two successive trusses), growth continues on live substems until their own trusses stop them too; '
            'declared axes only, generated trusses are model events (user_stated applicability), nothing is a measurement')


def r10(x: float) -> float:
    return float(f'{x:.{SIGNIFICANT_DIGITS}g}')


# ---------------------------------------------------------------------------------------------------------------
# Scenario
# ---------------------------------------------------------------------------------------------------------------

def default_start_time(route: str, days_at_start_value: int = 30, today: datetime | None = None) -> str:
    """Seed route: the declared sowing time is `days_at_start` days before today's 06:00 (+09:00), so the forecast
    origin (boundary day 30) is today. Fixture route: the observation boundary is today's 06:00."""
    today = today or datetime.now(tz=aware('2026-01-01T00:00:00+09:00').tzinfo).replace(hour=6, minute=0, second=0, microsecond=0)
    start = today - timedelta(days=days_at_start_value) if route == 'seed_route' else today
    return start.isoformat()


def default_scenario(route: str, horizons_days: list[int] | None = None, members: int = 4, seed: int = 11, rays: int = QUANTITATIVE_RAYS,
                     start_time: str | None = None, source_run: str | None = None, fixture_day: int = 30, scenario_id: str | None = None,
                     p1_package: str | None = None, state_source: str = 'fixture_truth', sowing_datetime: str | None = None,
                     diagnostic: bool = False, rerun_trigger: str = 'initial_issue', base_only: bool = False) -> dict:
    """Declared scenario; climate values come from the registry's environment rows (not typed here)."""
    require(route in ROUTES, f'route must be one of {ROUTES}')
    engine = ROUTE_ENGINE[route]
    if engine == 'pipeline':
        require(bool(p1_package), 'p1_pipeline route needs a P1 package path (relative to the project root)')
    start_time = start_time or (default_start_time(route) if engine == 'v3_direct' else None)
    return dict(schema=SCENARIO_SCHEMA, scenario_id=scenario_id or f'{route}_default', route=route, engine=engine,
                engine_scope='multi_instance_canopy_placement' if engine == 'pipeline' else 'single_plant_no_canopy',
                start_time=start_time,
                start_time_status=('taken_from_P1_captures_at_prepare' if engine == 'pipeline' else
                                   'declared_assumption; real sowing date unrecorded (seed_route) / observation boundary (stage2_fixture)'),
                horizons_days=horizons_days or [1, 3, 7, 14], lamps=DEFAULT_LAMPS, light_input=DEFAULT_LIGHT,
                ratio_pointer=DEFAULT_POINTER, lights_on_clock_time=None, rays=rays,
                issue_class='diagnostic' if (diagnostic or rays < QUANTITATIVE_RAYS) else 'quantitative',
                ensemble=dict(size=members, seed=seed), source_run=source_run, fixture_day=fixture_day,
                p1_package=p1_package, state_source=state_source, axis_hidden_state=dict(AXIS_HIDDEN_DEFAULT),
                sowing_datetime=sowing_datetime, sowing_source='user_stated' if sowing_datetime else None,
                rerun_trigger=rerun_trigger, scenario_set_mode='base_only' if base_only else 'all',
                declared_assumptions=['start = sowing + 30 d (user decision 2026-10-11; real sowing dates unrecorded) unless --sowing-datetime is given',
                                      'reference PPFD is user_assumed until the meter arrives; absolute outputs conditional',
                                      'lamp placement is a declared guess (uniform point lamp above the trays)'])


def validate_scenario(s: dict) -> dict:
    require(s.get('schema') == SCENARIO_SCHEMA, 'unknown scenario schema')
    require(s['route'] in ROUTES and s.get('engine') == ROUTE_ENGINE[s['route']], 'route/engine mismatch')
    if s['engine'] == 'v3_direct':
        aware(s['start_time'])
    h = s['horizons_days']
    require(isinstance(h, list) and h and all(type(x) is int and x > 0 for x in h) and len(set(h)) == len(h), 'horizons: unique positive integer days')
    require(type(s['rays']) is int and s['rays'] >= 1, 'rays must be a positive integer')
    require(s.get('issue_class') in ISSUE_CLASSES, 'issue_class must be quantitative or diagnostic')
    require(s['rays'] >= QUANTITATIVE_RAYS or s['issue_class'] == 'diagnostic',
            f'fewer than {QUANTITATIVE_RAYS} rays requires issue_class = diagnostic (--rays N --diagnostic)')
    e = s['ensemble']
    require(type(e['size']) is int and e['size'] >= 1 and type(e['seed']) is int, 'ensemble size/seed must be integers')
    require(isinstance(s['lamps'], list) and s['lamps'], 'at least one lamp')
    if s['route'] == 'stage2_fixture':
        require(bool(s.get('source_run')), 'stage2_fixture needs source_run (a verified V3 run directory)')
        require(type(s.get('fixture_day')) is int and s['fixture_day'] >= 1, 'fixture_day must be a positive integer')
    if s['engine'] == 'pipeline':
        require(bool(s.get('p1_package')) and (ROOT / s['p1_package']).is_dir(), 'p1_package must be an existing directory under the project root')
        require(s.get('state_source') in ('main', 'fixture_truth'), 'state_source must be main or fixture_truth')
    if s.get('lights_on_clock_time') is not None:
        datetime.strptime(s['lights_on_clock_time'], '%H:%M')
    if s.get('sowing_datetime') is not None:
        aware(s['sowing_datetime'])
        require(s.get('sowing_source') == 'user_stated', 'a given sowing_datetime is user_stated')
    s.setdefault('rerun_trigger', 'initial_issue')  # scenarios written before F9
    require(s['rerun_trigger'] in RERUN_TRIGGERS, f'rerun_trigger must be one of {sorted(RERUN_TRIGGERS)}')
    s.setdefault('scenario_set_mode', 'all')  # F10: base_only = diagnostic issue of the reference scenario alone (never pooled either way)
    require(s['scenario_set_mode'] in ('all', 'base_only'), 'scenario_set_mode must be all or base_only')
    return s


def environment_values(overrides: dict) -> dict:
    env = {k.split(':', 1)[1]: v for k, v in overrides['environment'].items()}
    for key in ('temp_day_c', 'temp_night_c', 'light_hours', 'co2_ppm', 'rh_pct'):
        require(key in env, f'registry lacks environment row for {key}')
    return env


def days_at_start(overrides: dict) -> int:
    row = overrides['layer_only'].get('DAYS_AT_START')
    require(row is not None and type(row['value']) is int and row['value'] >= 0, 'registry needs DAYS_AT_START')
    return row['value']


def time_since_lights_on_min(start_time: str, lights_on: str | None) -> int | str:
    if lights_on is None:
        return 'unavailable'
    boundary = aware(start_time)
    on = datetime.strptime(lights_on, '%H:%M')
    return ((boundary.hour * 60 + boundary.minute) - (on.hour * 60 + on.minute)) % 1440


def os_relpath(target: Path, start: Path) -> str:
    return os.path.relpath(str(target), str(start))


# ---------------------------------------------------------------------------------------------------------------
# B5 — sowing provenance
# ---------------------------------------------------------------------------------------------------------------

def registry_sowing(xlsx: Path = REGISTRY_XLSX) -> dict[str, str]:
    """Read-only look for a per-plant sowing column in the human plant registry (개체목록 sheet). Absent -> {}."""
    if not Path(xlsx).exists():
        return {}
    try:
        import openpyxl
    except ImportError:
        return {}
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    out: dict[str, str] = {}
    for ws in wb.worksheets:
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        col = next((i for i, h in enumerate(header) if h in SOWING_COLUMNS), None)
        pid = next((i for i, h in enumerate(header) if h == 'plant_id'), None)
        pid = next((i for i, h in enumerate(header) if h == 'slot_id'), None) if pid is None else pid
        if col is None or pid is None:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[col] not in (None, '') and row[pid] not in (None, ''):
                value = row[col]
                out[str(row[pid])] = value.isoformat() if hasattr(value, 'isoformat') else str(value)
    return out


def resolve_sowing(scenario: dict, d0: int, registry_xlsx: Path = REGISTRY_XLSX) -> dict:
    """sowing_datetime + sowing_source: user_stated > plant registry column > assumed (seed route) > unavailable."""
    if scenario.get('sowing_datetime'):
        return dict(sowing_datetime=aware(scenario['sowing_datetime']).isoformat(), sowing_source='user_stated', per_plant={})
    per_plant = registry_sowing(registry_xlsx)
    if per_plant:
        return dict(sowing_datetime='per_plant', sowing_source='plant_registry_user_record', per_plant=per_plant)
    if scenario['route'] == 'seed_route':
        return dict(sowing_datetime=aware(scenario['start_time']).isoformat(), sowing_source='assumed_issue_minus_days_at_start',
                    per_plant={}, note=f'seed route start_time = issue boundary minus DAYS_AT_START ({d0} d); real sowing dates unrecorded')
    return dict(sowing_datetime='unavailable', sowing_source='unavailable', per_plant={},
                note='no --sowing-datetime, no sowing column in the plant registry (개체목록 has 이식일 only)')


def days_after_sowing_fn(sowing: dict, boundary_fn):
    def f(h: int, pid: str | None = None):
        when = sowing['per_plant'].get(pid) if sowing['sowing_datetime'] == 'per_plant' else sowing['sowing_datetime']
        if not when or when == 'unavailable':
            return 'unavailable'
        value = r10((boundary_fn(h) - aware(when)).total_seconds() / 86400)
        return int(value) if value == int(value) else value
    return f


# ---------------------------------------------------------------------------------------------------------------
# Members — v3_direct engine
# ---------------------------------------------------------------------------------------------------------------

def _write_common_inputs(inputs: Path, scenario: dict, overrides: dict, intervals: int, member_seed: int) -> dict:
    env = environment_values(overrides)
    co2 = overrides['environment']['environment.csv:co2_ppm']
    climate = make_climate_csv(inputs, scenario['start_time'], intervals, env['temp_day_c']['value'], env['temp_night_c']['value'],
                               env['light_hours']['value'], env['rh_pct']['value'], co2['value'],
                               source_kind=env['temp_day_c']['source_kind'], co2_source_kind=co2['source_kind'],
                               t_base_c=overrides['parameters']['BASE_C']['value'])
    write_csv(inputs / 'lamps.csv', [dict(l) for l in scenario['lamps']], ['lamp_id', 'x_m', 'y_m', 'z_m', 'power_umol_s', 'source_kind'])
    write_csv(inputs / 'light_input.csv', [dict(scenario['light_input'])], ['reference_ppfd_umol_m2_s', 'source_kind', 'measured_at'])
    target = (ROOT / scenario['ratio_pointer']).resolve()
    require(target.is_file(), f'ratio pointer target missing: {target}')
    (inputs / 'light_ratio_latest.txt').write_text(Path(os_relpath(target, inputs)).as_posix() + '\n', encoding='utf-8')
    if (inputs / 'ratios.csv').exists():
        (inputs / 'ratios.csv').unlink()  # demo's uniform placeholder; the human slot file is the pointer target
    pointer = resolve_ratio_pointer(inputs / 'light_ratio_latest.txt')
    req = read_json(inputs / 'request.json')
    req['runtime']['seed'] = member_seed
    req['runtime']['rays'] = scenario['rays']
    return dict(climate=climate, pointer=pointer, request=req)


def build_member_inputs(folder: Path, scenario: dict, registry: list[dict], index: int, member_values: dict, issue_boundary: int) -> dict:
    """Create one frozen V3 input package for member `index` under `folder/members/mNN/inputs` (v3_direct engine)."""
    from tools.vtc_development import demo
    overrides = build_request_overrides(registry, member_values)
    inputs = folder / 'members' / f'm{index:02d}' / 'inputs'
    run_id = f'{folder.name}_m{index:02d}'
    max_h = max(scenario['horizons_days'])
    if scenario['route'] == 'seed_route':
        intervals = issue_boundary + max_h
        demo(inputs, intervals=intervals, rays=scenario['rays'], scenario='sufficient', seed=scenario['ensemble']['seed'])
        common = _write_common_inputs(inputs, scenario, overrides, intervals, scenario['ensemble']['seed'])
        req = common['request']
    else:
        from tools.vtc_development_observation import fixture_from_native
        intervals = max_h
        fixture_from_native(ROOT / scenario['source_run'], inputs, scenario['fixture_day'], intervals)
        common = _write_common_inputs(inputs, scenario, overrides, intervals, scenario['ensemble']['seed'])
        req = common['request']
        # The imported inventory fixes PAIRS; the registry row is kept at the source run's value for this route.
        source_pairs = req['parameters']['PAIRS']['value']
        overrides['parameters']['PAIRS'] = dict(overrides['parameters']['PAIRS'], value=source_pairs)
        member_values = {k: v for k, v in member_values.items() if k != 'NR_LEAFLET_PAIRS_BY_RANK'}
    req['start_time'] = aware(scenario['start_time']).isoformat()
    req['run_id'] = run_id
    req['scenario_id'] = scenario['scenario_id']
    req['forecast_layer'] = dict(forecast_run_id=folder.name, member_index=index, member_values=member_values,
                                 inactive_in_engine=[i['key'] for i in overrides['inactive_in_engine']],
                                 engine_default_retained=[i['engine_key'] for i in overrides['engine_default_retained']])
    req = apply_overrides(req, overrides)
    write_json(inputs / 'request.json', req)
    from tools.vtc_development_io import load_package
    pkg = load_package(inputs)  # V3's own contract check; raises ContractError on any violation
    return dict(index=index, inputs=inputs, run_id=run_id, member_values=member_values, overrides=overrides, request=req,
                intervals=intervals, pointer=common['pointer'], climate=common['climate'], source_hashes=pkg['source_hashes'])


def run_member(inputs: Path, run_dir: Path, java: str = 'java', timeout_s: float = 1200) -> Path:
    """prepare -> run -> collect through V3's public functions; returns the published results folder."""
    from tools.vtc_development import collect, prepare
    from tools.vtc_growth_bridge import run
    prepare(inputs, run_dir)
    execution = run(run_dir, java=java, timeout_s=timeout_s)
    require(execution.get('status') == 'verified', 'native execution not verified; forecast member rejected')
    collect(run_dir)
    return run_dir / 'results'


def load_member_results(results: Path, request: dict) -> dict:
    states: dict[int, dict] = defaultdict(dict)
    for s in rows(results / 'states.csv'):
        states[int(s['day'])][s['organ_id']] = s
    events = rows(results / 'events.csv')
    traits: dict[int, dict] = defaultdict(dict)
    if (results / 'compound_traits.csv').exists():
        for t in rows(results / 'compound_traits.csv'):
            traits[int(t['day'])][t['organ_id']] = t
    manifest = read_json(results / 'manifest.json')
    require(manifest.get('run_status') == RUN_STATUS and manifest.get('software_status') == 'verified_native_trajectory', 'member results not verified')
    return dict(states=states, events=events, traits=traits, manifest=manifest, request=request,
                plants=[p['plant_instance_id'] for p in rows(results.parent / 'inputs' / 'plants.csv')], checkpoints=[])


# ---------------------------------------------------------------------------------------------------------------
# Members — pipeline engine (B1)
# ---------------------------------------------------------------------------------------------------------------

def pipeline_scenario(folder: Path, scenario: dict, overrides: dict, intervals: int, seed: int) -> Path:
    """One pipeline scenario via Codex's public `scenario_fixture`, then the member parameters and the registry climate/lamps."""
    from tools.tomato_fsp_acceptance import scenario_fixture
    path = scenario_fixture(ROOT / scenario['p1_package'], folder, intervals=intervals, rays=scenario['rays'], seed=seed,
                            state_source=scenario['state_source'])
    scen = read_json(path)
    scen['parameters'] = apply_overrides(dict(parameters=scen['parameters']), overrides)['parameters']
    env = environment_values(overrides)
    co2 = overrides['environment']['environment.csv:co2_ppm']
    scen['environment'] = [dict(light_hours=env['light_hours']['value'], temp_day_c=env['temp_day_c']['value'], temp_night_c=env['temp_night_c']['value'],
                                co2_ppm=co2['value'], source_kind=env['temp_day_c']['source_kind']) for _ in range(intervals)]
    scen['lamps'] = [dict(l) for l in scenario['lamps']]
    package = ROOT / scenario['p1_package']
    if ax.has_axes(package):
        from tools.fsp_axes import select_axes
        vertices = ax.selected_vertices(package, scen['captures_by_instance'], scenario['state_source'])
        thermal = dict(GERMINATION_DD=overrides['parameters']['GERMINATION_DD']['value'], PHYLLOCHRON_DD=overrides['parameters']['PHYLLOCHRON_DD']['value'])
        records, notes = ax.build_axis_records(package, scen['captures_by_instance'], vertices, overrides['axis_parameters'],
                                               scenario.get('axis_hidden_state', AXIS_HIDDEN_DEFAULT), thermal)
        ids = {f'{v["plant_instance_id"]}|{v["vertex_id"]}' for v in vertices}
        scen['organ_initialization'] = {k: v for k, v in scen['organ_initialization'].items() if k in ids}
        supplemented = ax.truss_geometry_supplement(scen['organ_initialization'], vertices)
        scen.update(axis_contract=ax.AXIS_CONTRACT, topology_scope=TOPOLOGY_SCOPE, topology_completeness='unavailable', axes=records,
                    nutrient_ledger_required=False, allow_conditional_surfaces=True)
        scen.pop('nutrient_scenarios', None)
        selected, membership = select_axes(package, scen, vertices)  # engine-owner validation (public); fills truss_history / truss_rule_status
        scen['axes'] = selected
        write_json(folder / 'axis_notes.json', dict(notes=notes, truss_supplement_ids=supplemented, truss_geometry_status=ax.ZERO_EXTENT_LABEL,
                                                     axes=[dict(axis_id=a['axis_id'], plant_instance_id=a['plant_instance_id'], axil_parent_rank=a['parent_rank'],
                                                                segment_index=a.get('segment_index', 'unavailable'), truss_rule_status=a['truss_rule_status'],
                                                                truss_history=a['truss_history']) for a in selected]))
    write_json(path, scen)
    return path


def build_pipeline_member(folder: Path, scenario: dict, registry: list[dict], index: int, member_values: dict, timeout_s: float,
                          scenario_values: dict | None = None, label: str | None = None, scenario_edit=None) -> dict:
    """prepare_run -> run (which commits checkpoints) through `tools.tomato_fsp_pipeline`; never `advance_interval`.

    `scenario_edit` (F9, stress fixtures only) is a function applied to the engine scenario dict before `prepare_run`; it must
    return a dict whose `initialization_provenance` carries `NO_REUSE_MARK`, so the run can never be consumed as a forecast member.
    """
    from tools.tomato_fsp_pipeline import prepare_run, run
    overrides = build_request_overrides(registry, member_values, scenario_values)
    mdir = (folder / 'scenarios' / label if label else folder) / 'members' / f'm{index:02d}'
    intervals = max(scenario['horizons_days'])
    scen_path = pipeline_scenario(mdir / 'scenario', scenario, overrides, intervals, scenario['ensemble']['seed'])
    scen = read_json(scen_path)
    if scenario_edit is not None:
        scen = scenario_edit(scen)
        require(NO_REUSE_MARK in scen.get('initialization_provenance', ''), 'an edited (stress) engine scenario must carry the no-reuse mark')
        write_json(scen_path, scen)
    run_dir = mdir / 'run'
    prepared = prepare_run(ROOT / scenario['p1_package'], scen_path, run_dir)
    collection = run(run_dir, timeout=timeout_s)
    require(collection.get('status') == 'complete_conditional_trajectory' and collection['last_valid_boundary'] == intervals,
            'pipeline trajectory incomplete; forecast member rejected')
    return dict(index=index, run_dir=run_dir, scenario_path=scen_path, scenario=scen, prepared=prepared, collection=collection,
                member_values=member_values, overrides=overrides, intervals=intervals, scenario_values=scenario_values or {})


MAX_MEMBER_RETRIES = 2


def build_pipeline_member_with_retry(folder: Path, scenario: dict, registry: list[dict], index: int, member_values: dict, timeout_s: float,
                                     scenario_values: dict | None, label: str | None, transients: list[dict], scenario_edit=None) -> dict:
    """Rebuild a member from scratch (fresh folder) when the engine's atomic folder rename is denied by the OS/OneDrive.

    Only `PermissionError` (WinError 5) during publication is retried; the failed attempt folder is moved aside as
    `mNN_failed_attempt_k` (never reused, never advanced), and every retry is recorded for the manifest. Any other error propagates.
    """
    import shutil
    for attempt in range(MAX_MEMBER_RETRIES + 1):
        try:
            return build_pipeline_member(folder, scenario, registry, index, member_values, timeout_s, scenario_values, label, scenario_edit)
        except PermissionError as exc:
            mdir = (folder / 'scenarios' / label if label else folder) / 'members' / f'm{index:02d}'
            aside = mdir.parent / f'm{index:02d}_failed_attempt_{attempt + 1}'
            transients.append(dict(scenario=label or 'base', member=index, attempt=attempt + 1, error=f'{type(exc).__name__}: {exc}', moved_to=str(aside.relative_to(folder)), issued_at=now_iso()))
            if attempt == MAX_MEMBER_RETRIES:
                raise
            try:
                shutil.move(str(mdir), str(aside))
            except OSError:
                shutil.move(str(mdir), str(aside) + '_' + uuid.uuid4().hex[:6])
    raise ContractError('unreachable')


def load_pipeline_member(built: dict) -> dict:
    """States/events from the frozen initial state and the committed checkpoints only."""
    run_dir: Path = built['run_dir']
    states: dict[int, dict] = defaultdict(dict)
    for s in rows(run_dir / 'native_inputs' / 'initial_state.csv'):
        states[0][s['organ_id']] = s
    events: list[dict] = []
    checkpoints = []
    for k in range(1, built['intervals'] + 1):
        step = run_dir / 'committed' / f'step_{k:04d}'
        meta = read_json(step / 'checkpoint.json')
        require(meta['status'] == 'committed_once' and meta['native_advance_count'] == 1, 'checkpoint not committed once')
        for name, h in meta['files'].items():
            require(sha(step / name) == h, f'committed file changed: {name}')
        for s in rows(step / 'state.csv'):
            require(int(s['day']) == k, 'checkpoint state day mismatch')
            states[k][s['organ_id']] = s
        if (step / 'development_events.csv').stat().st_size > 0:
            events.extend(rows(step / 'development_events.csv'))
        checkpoints.append(dict(boundary=k, checkpoint_sha256=sha(step / 'checkpoint.json'), cache_key=meta['cache_key'], state_time=meta['state_time']))
    scen = built['scenario']
    request = dict(parameters=scen['parameters'], frame='plant_local_mm', mode='vtc_development_v3', start_time=scen['start_time'],
                   runtime=dict(rays=scen['rays'], seed=scen['seed']), initialization_route='p1_conditioned', intervals=built['intervals'],
                   axis_contract=scen.get('axis_contract'), axes=scen.get('axes', []))
    axis_inventory = {k: ax.committed_axis_inventory(run_dir, k) for k in range(1, built['intervals'] + 1)} if scen.get('axis_contract') else {}
    return dict(states=states, events=events, traits={}, manifest=built['collection'], request=request,
                plants=sorted(scen['captures_by_instance']), checkpoints=checkpoints, axis_inventory=axis_inventory,
                axis_sums=ax.axis_sums_check(run_dir) if scen.get('axis_contract') else None,
                run_dir=run_dir, intervals=built['intervals'])


# ---------------------------------------------------------------------------------------------------------------
# F10 — per-interval carbon budget trajectory of pipeline members (does the allocation cap ever bind?)
# ---------------------------------------------------------------------------------------------------------------
BIND_TOL_GC = 1e-10


def carbon_trajectory(run_id: str, scenario_id: str, scenario_value, member_index: int, mem: dict, rays: int) -> list[dict]:
    """One row per committed day and plant from `carbon_ledger.csv` + `axis_flux.csv`; `cap_binding` = allocated < demand.

    The engine law (vtc_development_native.rgg): budget = reserve_start + gross - maintenance; funded = min(max(0, budget), demand).
    Turning points are local extrema of reserve_end (sign change of the daily change). Events listed: determinate stops and truss births.
    """
    run_dir: Path = mem['run_dir']
    rows_out: list[dict] = []
    for k in range(1, mem['intervals'] + 1):
        step = run_dir / 'committed' / f'step_{k:04d}'
        ledger = {r['plant_instance_id']: r for r in rows(step / 'carbon_ledger.csv')}
        flux = rows(step / 'axis_flux.csv') if (step / 'axis_flux.csv').is_file() else []
        for pid, led in sorted(ledger.items()):
            demand = sum(float(r['demand_gC']) for r in flux if r['plant_instance_id'] == pid)
            allocated = sum(float(r['allocated_gC']) for r in flux if r['plant_instance_id'] == pid)
            budget = float(led['reserve_start_gC']) + float(led['gross_gC']) - float(led['maintenance_gC'])
            evs = sorted(f'{e["event"]}:{e["organ_id"]}' for e in mem['events'] if int(e['day']) == k
                         and (e['event'] == 'determinate_stop' or (e['event'] == 'funded_birth' and 'Truss' in e['organ_id'])))
            rows_out.append(dict(forecast_run_id=run_id, scenario_id=scenario_id, scenario_value=scenario_value, member_index=member_index, day=k, plant_instance_id=pid,
                                 rays=rays, reserve_start_gC=float(led['reserve_start_gC']), gross_gC=float(led['gross_gC']), maintenance_gC=float(led['maintenance_gC']),
                                 budget_gC=budget, demand_gC=demand, allocated_gC=allocated, funded_fraction=(allocated / demand) if demand > 0 else 'n/a',
                                 cap_binding=allocated < demand - BIND_TOL_GC, reserve_end_gC=float(led['reserve_end_gC']), deficit_gC=float(led['deficit_gC']),
                                 residual_gC=float(led['residual_gC']), turning_point='', events_this_day=';'.join(evs)))
    for pid in {r['plant_instance_id'] for r in rows_out}:
        seq = [r for r in rows_out if r['plant_instance_id'] == pid]
        for a, b, c in zip(seq, seq[1:], seq[2:]):
            d1, d2 = b['reserve_end_gC'] - a['reserve_end_gC'], c['reserve_end_gC'] - b['reserve_end_gC']
            if d1 > 0 and d2 < 0:
                b['turning_point'] = 'peak'
            elif d1 < 0 and d2 > 0:
                b['turning_point'] = 'trough'
    return rows_out


def cap_summary(trajectory: list[dict], members_by_scenario: dict[str, list[dict]], horizon_days: int) -> dict:
    """Per scenario: first binding day (or none within the horizon), turning points, growth-ending events before binding."""
    out: dict[str, Any] = {}
    for label, mems in members_by_scenario.items():
        per_member = []
        for i, mem in enumerate(mems):
            mine = [r for r in trajectory if r['scenario_id'] == label and r['member_index'] == i]
            binding_days = sorted({r['day'] for r in mine if r['cap_binding']})
            first = binding_days[0] if binding_days else None
            limit = first if first is not None else horizon_days
            stops = sorted((int(e['day']), e.get('axis_id', 'M'), e['organ_id']) for e in mem['events'] if e['event'] == 'determinate_stop' and int(e['day']) <= limit)
            trusses = sorted((int(e['day']), e.get('axis_id', 'M'), e['organ_id']) for e in mem['events'] if e['event'] == 'funded_birth' and 'Truss' in e['organ_id'] and int(e['day']) <= limit)
            axes = [a['axis_id'] for a in mem['request'].get('axes', [])] or ['M']
            stopped_axes = {a for _, a, _ in stops}
            per_member.append(dict(member_index=i, first_binding_day=first, binding_days=binding_days,
                                   turning_points=[dict(day=r['day'], plant_instance_id=r['plant_instance_id'], kind=r['turning_point'], reserve_end_gC=r['reserve_end_gC']) for r in mine if r['turning_point']],
                                   reserve_end_last_gC={r['plant_instance_id']: r['reserve_end_gC'] for r in mine if r['day'] == horizon_days},
                                   determinate_stops_before_binding=[dict(day=d, axis_id=a, organ_id=o) for d, a, o in stops],
                                   truss_births_before_binding=[dict(day=d, axis_id=a, organ_id=o) for d, a, o in trusses],
                                   all_declared_axes_stopped_before_binding=set(axes) <= stopped_axes))
        firsts = [m['first_binding_day'] for m in per_member]
        out[label] = dict(binds_within_horizon=any(f is not None for f in firsts), first_binding_day_by_member=firsts,
                          first_binding_day_min=min((f for f in firsts if f is not None), default=None), members=per_member)
    return dict(horizon_days=horizon_days, rays=trajectory[0]['rays'] if trajectory else None, per_scenario=out,
                law='budget = reserve_start + gross - maintenance; funded = min(max(0, budget), demand); the cap binds when allocated < demand',
                label='binding day = property of this fixture\'s declared state (model_property), not a Micro-Tom prediction; assumption error accumulates with horizon')


# ---------------------------------------------------------------------------------------------------------------
# B3 — eligibility inputs per route
# ---------------------------------------------------------------------------------------------------------------

def controls_from_native_run(run_dir: Path, day: int) -> dict[str, dict]:
    """Apex counters at `day` of a verified V3 run (same reading `fixture_from_native` uses), per plant."""
    states = [s for s in rows(run_dir / 'results' / 'states.csv') if int(s['day']) == day]
    apices = [a for a in rows(run_dir / 'results' / 'apices.csv') if int(a['day']) == day]
    out = {}
    for pid in sorted({s['plant_instance_id'] for s in states}):
        ps = [s for s in states if s['plant_instance_id'] == pid]
        ids = {s['organ_id'] for s in ps}
        active = next((a for a in apices if a['organ_id'] in ids and a['stopped'] == 'false'), None)
        stopped = active is None
        if active is None:
            active = [a for a in apices if a['organ_id'] in ids][-1]
        axis = next(s for s in ps if s['organ_id'] == active['organ_id'])['axis_id']
        out[pid] = dict(thermal_dd=float(active['plant_thermal_dd']), produced=int(active['produced_events']),
                        leaf_order=max(int(s['true_leaf_order']) for s in ps), truss_count=sum(s['class'] == 'TrussAxis' for s in ps),
                        axis_leaf_count=sum(s['class'] == 'CompoundLeaf' and s['axis_id'] == axis for s in ps), stopped=stopped)
    return out


def eligibility_inputs(scenario: dict, folder: Path, overrides: dict, d0: int) -> tuple[dict[str, dict], int, int]:
    """(controls per plant, start_interval, days) for `microtom_parameters.eligibility`."""
    max_h = max(scenario['horizons_days'])
    if scenario['route'] == 'seed_route':
        return {'seed': dict(thermal_dd=0.0, produced=0, leaf_order=0, truss_count=0, axis_leaf_count=0, stopped=False)}, 0, d0 + max_h
    if scenario['route'] == 'stage2_fixture':
        return controls_from_native_run(ROOT / scenario['source_run'], scenario['fixture_day']), 0, max_h
    probe = pipeline_scenario(folder / 'eligibility_probe', scenario, overrides, max_h, scenario['ensemble']['seed'])
    scen = read_json(probe)
    if scen.get('axis_contract'):
        return ax.axis_controls(scen['axes']), 0, max_h
    plants = scen['plant_initialization']
    return {pid: dict(c, stopped=False) for pid, c in plants.items()}, 0, max_h


def environment_rows(overrides: dict, days: int) -> list[dict]:
    env = environment_values(overrides)
    return [dict(light_hours=env['light_hours']['value'], temp_day_c=env['temp_day_c']['value'], temp_night_c=env['temp_night_c']['value']) for _ in range(days)]


def plain_parameters(overrides: dict) -> dict:
    from tools.vtc_development import parameters
    base = {k: v['value'] for k, v in parameters().items()}
    base.update({k: v['value'] for k, v in overrides['parameters'].items()})
    return base


# ---------------------------------------------------------------------------------------------------------------
# Aggregation
# ---------------------------------------------------------------------------------------------------------------

def _ancestor_at(states_day0: dict, states_day: dict, organ_id: str) -> str:
    cur = organ_id
    while cur:
        if cur in states_day0:
            return cur
        cur = states_day[cur]['parent_id']
    return ''


def _quantiles(values: list[float]) -> tuple[float, float, float]:
    return r10(statistics.median(values)), r10(min(values)), r10(max(values))


def _descendant_length_mm(states_day: dict, compound_id: str) -> float:
    children = defaultdict(list)
    for oid, s in states_day.items():
        children[s['parent_id']].append(oid)
    total = 0.0
    todo = list(children[compound_id])
    while todo:
        oid = todo.pop()
        s = states_day[oid]
        if s['class'] in ('Petiole', 'RachisSegment'):
            total += float(s['length_m']) * 1e3
        todo.extend(children[oid])
    return total


def plant_metrics_from_states(states_day: dict, pid: str) -> dict[str, float]:
    ps = [s for s in states_day.values() if s['plant_instance_id'] == pid]
    base_z = min((float(s['z_m']) for s in ps if s['class'] == 'Plant'), default=0.0)
    tips = [max(float(s['z_m']), float(s['tip_z_m'])) for s in ps if s['class'] == 'Leaflet']
    return dict(stem_length_mm=sum(float(s['length_m']) for s in ps if s['class'] == 'Internode') * 1e3,
                height_top_leaf_axis_mm=(max(tips) - base_z) * 1e3 if tips else 0.0,
                structural_mass_gDM=sum(float(s['mass_gDM']) for s in ps), root_mass_gDM=sum(float(s['mass_gDM']) for s in ps if s['class'] == 'Root'),
                compounds=float(sum(s['class'] == 'CompoundLeaf' for s in ps)), trusses=float(sum(s['class'] == 'TrussAxis' for s in ps)))


def probability_basis(event_type: str, members: list[dict], axis_key: tuple[str, str] | None = None) -> str:
    """B4: rule_deterministic when every member shares the parameters that produce the event; else scenario_fraction.

    With `axis_key = (plant, axis)` the axis's own fsp-axes-1 parameters (APPLY_DETERMINATE_STOP) join the governing set."""
    keys: tuple[str, ...] = THERMAL_RULE_KEYS
    if event_type.startswith('organ_initiation:') and event_type.split(':')[1] in PAIR_CLASSES:
        keys = (*keys, 'PAIRS')
    values = []
    for m in members:
        v = [json.dumps(m['request']['parameters'][k]['value'], sort_keys=True) for k in keys if k in m['request']['parameters']]
        if axis_key is not None:
            rec = next((a for a in m['request'].get('axes', []) if (a['plant_instance_id'], a['axis_id']) == axis_key), None)
            v += [json.dumps(rec['parameters'][k]['value']) for k in AXIS_RULE_KEYS] if rec else ['axis_absent']
        values.append(tuple(v))
    if not values:
        return 'unavailable'
    return 'rule_deterministic' if all(v == values[0] for v in values) else 'scenario_fraction'


def _axis_of(state: dict, axis_mode: bool) -> str:
    return state.get('axis_id', 'M') if axis_mode else 'M'


def aggregate(run_id: str, members: list[dict], horizons: list[int], issue_boundary: int, das, frame: str,
              scenario_id: str | None = None, scenario_value: Any = 'n/a') -> dict[str, list[dict]]:
    """Members' states/events -> forecast_mtg, forecast_intervals, forecast_events (own id namespace).

    Axis columns: `axis_id` (logical M/S<k> in fsp-axes-1 members, 'M' legacy label otherwise), `axil_parent_rank` (the engine
    state's parent_rank = side attachment rank, 0 for M), `organ_segment_index` (immutable, state.segment_index),
    `axis_segment_index` (current counter from the committed axis inventory at that boundary), `truss_provenance`."""
    n = len(members)
    axis_mode = all(m['request'].get('axis_contract') for m in members)
    mtg, intervals, events = [], [], []
    source_kind = 'user_assumed'  # conditional simulation under user_assumed / literature_other_cultivar parameters
    fid_of = lambda key: forecast_identifier(run_id, key, scenario_id)
    tag = dict(scenario_id=scenario_id or 'base', scenario_value=scenario_value)
    for h in horizons:
        day = issue_boundary + h
        per_organ: dict[tuple[str, str], list[float]] = defaultdict(list)
        organ_meta: dict[str, dict] = {}
        presence: dict[str, int] = defaultdict(int)
        new_organ_birth: dict[str, list[int]] = defaultdict(list)
        plant_vals: dict[tuple[str, str], list[float]] = defaultdict(list)
        plant_events: dict[tuple[str, str], list[int | None]] = defaultdict(list)
        plant_counts: dict[tuple[str, str], list[int]] = defaultdict(list)
        axis_vals: dict[tuple[str, str, str], list[float]] = defaultdict(list)       # (pid, axis, quantity|unit)
        axis_ev_first: dict[tuple[str, str, str], list[int | None]] = defaultdict(list)
        axis_ev_count: dict[tuple[str, str, str], list[int]] = defaultdict(list)
        for mi, m in enumerate(members):
            sd, s0 = m['states'][day], m['states'][issue_boundary]
            traits = m['traits'].get(day, {})
            inventory = m.get('axis_inventory', {}).get(day, {})
            for oid, s in sd.items():
                fid = fid_of(oid)
                presence[oid] += 1
                new = oid not in s0
                root = oid if not new else _ancestor_at(s0, sd, s['parent_id'])
                if new:
                    new_organ_birth[oid].append(int(s['birth_day']) - issue_boundary)
                axis_id = _axis_of(s, axis_mode)
                organ_meta.setdefault(oid, dict(cls=s['class'], pid=s['plant_instance_id'], axis=axis_id))
                inv = inventory.get((s['plant_instance_id'], axis_id), {})
                mtg.append(dict(forecast_run_id=run_id, **tag, trajectory_id=f'm{mi:02d}', forecast_id=fid, rooted_at_vertex_id=root,
                                axis_id=axis_id, axil_parent_rank=s.get('parent_rank', 'unavailable') if axis_mode else 'unavailable',
                                organ_segment_index=s.get('segment_index', 'unavailable') if axis_mode else 'unavailable',
                                axis_segment_index=inv.get('segment_index', 'unavailable'),
                                developmental_axis_id=s.get('developmental_axis_id', s.get('axis_id', '')) if axis_mode else s.get('axis_id', ''),
                                truss_provenance=ax.truss_provenance(s, s0),
                                root_relation='new_organ' if new else 'continuation', source_organ_id=oid, **{'class': s['class']},
                                parent_forecast_id=fid_of(s['parent_id']) if s['parent_id'] else '',
                                plant_instance_id=s['plant_instance_id'], true_leaf_order=s['true_leaf_order'], pair_index=s['pair_index'],
                                side=s['side'], horizon_day=h, boundary_day=day, days_after_sowing=das(h, s['plant_instance_id']), frame=frame,
                                base_x_mm=r10(float(s['x_m']) * 1e3), base_y_mm=r10(float(s['y_m']) * 1e3), base_z_mm=r10(float(s['z_m']) * 1e3),
                                tip_x_mm=r10(float(s['tip_x_m']) * 1e3), tip_y_mm=r10(float(s['tip_y_m']) * 1e3), tip_z_mm=r10(float(s['tip_z_m']) * 1e3),
                                length_mm=r10(float(s['length_m']) * 1e3), width_mm=r10(float(s['width_m']) * 1e3), area_mm2=r10(float(s['area_m2']) * 1e6),
                                area_basis='full_outline' if s['class'] == 'Leaflet' else 'n/a',
                                surface_role='individual_leaflet_polygon' if s['class'] == 'Leaflet' else 'structural_proxy' if float(s['length_m']) > 0 else 'none',
                                azimuth_deg=r10(float(s['azimuth_deg'])), elevation_deg=r10(float(s['elevation_deg'])),
                                rachis_bend_alpha_deg=r10(float(s['rachis_bend_alpha_deg'])), direction_mechanism='fixed_generation',
                                status=VALUE_STATUS, source_kind=source_kind))
                for quantity, column, factor, unit in ORGAN_QUANTITIES.get(s['class'], []):
                    per_organ[(oid, f'{quantity}|{unit}')].append(float(s[column]) * factor)
                if s['class'] == 'CompoundLeaf':
                    per_organ[(oid, 'compound_length_mm|mm')].append(_descendant_length_mm(sd, oid))
                    t = traits.get(oid)
                    if t:
                        for key in ('compound_insertion_angle_deg', 'rachis_curvature_deg', 'blade_azimuth_deg'):
                            if t[key] != 'unavailable':
                                per_organ[(oid, f'{key}|deg')].append(float(t[key]))
            for pid in m['plants']:
                pm = plant_metrics_from_states(sd, pid)
                for quantity, unit in PLANT_QUANTITIES:
                    plant_vals[(pid, f'{quantity}|{unit}')].append(pm[quantity])
                for ev, cls in EVENT_CLASSES.items():
                    births = sorted(int(s['birth_day']) for s in sd.values() if s['plant_instance_id'] == pid and s['class'] == cls
                                    and issue_boundary < int(s['birth_day']) <= day)
                    plant_counts[(pid, ev)].append(len(births))
                    plant_events[(pid, ev)].append(births[0] - issue_boundary if births else None)
                stops = sorted(int(e['day']) for e in m['events'] if e['event'] == 'determinate_stop' and issue_boundary < int(e['day']) <= day
                               and sd.get(e['organ_id'], {}).get('plant_instance_id') == pid)
                plant_counts[(pid, 'apex_stop')].append(len(stops))
                plant_events[(pid, 'apex_stop')].append(stops[0] - issue_boundary if stops else None)
                if axis_mode:
                    for a in m['request'].get('axes', []):
                        if a['plant_instance_id'] != pid:
                            continue
                        aid = a['axis_id']
                        organs = [t for t in sd.values() if t['plant_instance_id'] == pid and _axis_of(t, True) == aid]
                        axis_vals[(pid, aid, 'axis_leaf_count|count')].append(float(sum(t['class'] == 'CompoundLeaf' for t in organs)))
                        axis_vals[(pid, aid, 'axis_leaf_area_mm2|mm2')].append(sum(float(t['area_m2']) for t in organs if t['class'] == 'Leaflet') * 1e6)
                        axis_vals[(pid, aid, 'axis_stem_length_mm|mm')].append(sum(float(t['length_m']) for t in organs if t['class'] == 'Internode') * 1e3)
                        axis_vals[(pid, aid, 'axis_truss_count_generated|count')].append(float(sum(t['class'] in ax.TRUSS_CLASSES and t['organ_id'] not in s0 for t in organs)))
                        axis_vals[(pid, aid, 'axis_truss_count_registry|count')].append(float(sum(t['class'] in ax.TRUSS_CLASSES and t['organ_id'] in s0 for t in organs)))
                        for ev, cls in (('new_leaf', 'CompoundLeaf'), ('new_truss_generated', 'TrussAxis')):
                            births = sorted(int(t['birth_day']) for t in organs if t['class'] == cls and issue_boundary < int(t['birth_day']) <= day)
                            axis_ev_count[(pid, aid, ev)].append(len(births))
                            axis_ev_first[(pid, aid, ev)].append(births[0] - issue_boundary if births else None)
                        astops = sorted(int(e['day']) for e in m['events'] if e['event'] == 'determinate_stop' and issue_boundary < int(e['day']) <= day
                                        and e.get('axis_id') == aid and sd.get(e['organ_id'], {}).get('plant_instance_id') == pid)
                        axis_ev_count[(pid, aid, 'apex_stop')].append(len(astops))
                        axis_ev_first[(pid, aid, 'apex_stop')].append(astops[0] - issue_boundary if astops else None)
        for (oid, qu), values in sorted(per_organ.items()):
            quantity, unit = qu.split('|')
            point, low, high = _quantiles(values)
            pid = organ_meta[oid]['pid']
            intervals.append(dict(forecast_run_id=run_id, **tag, forecast_id=fid_of(oid), plant_instance_id=pid, axis_id=organ_meta[oid]['axis'],
                                  quantity=quantity, unit=unit, horizon_day=h, days_after_sowing=das(h, pid), point=point, interval_low=low, interval_high=high,
                                  nominal_level=f'envelope_of_{len(values)}_of_{n}_members', interval_kind='declared_scenario_envelope', calibrated=False,
                                  n_members=n, n_members_present=presence[oid], presence_probability=r10(presence[oid] / n), status=VALUE_STATUS,
                                  source_kind=source_kind))
        for (pid, qu), values in sorted(plant_vals.items()):
            quantity, unit = qu.split('|')
            point, low, high = _quantiles(values)
            intervals.append(dict(forecast_run_id=run_id, **tag, forecast_id=fid_of(f'{pid}/Plant'), plant_instance_id=pid, axis_id='plant', quantity=quantity,
                                  unit=unit, horizon_day=h, days_after_sowing=das(h, pid), point=point, interval_low=low, interval_high=high,
                                  nominal_level=f'envelope_of_{len(values)}_of_{n}_members', interval_kind='declared_scenario_envelope', calibrated=False,
                                  n_members=n, n_members_present=len(values), presence_probability=1.0, status=VALUE_STATUS, source_kind=source_kind))
        for (pid, aid, qu), values in sorted(axis_vals.items()):
            quantity, unit = qu.split('|')
            point, low, high = _quantiles(values)
            intervals.append(dict(forecast_run_id=run_id, **tag, forecast_id=fid_of(f'{pid}/{aid}/Axis'), plant_instance_id=pid, axis_id=aid, quantity=quantity,
                                  unit=unit, horizon_day=h, days_after_sowing=das(h, pid), point=point, interval_low=low, interval_high=high,
                                  nominal_level=f'envelope_of_{len(values)}_of_{n}_members', interval_kind='declared_scenario_envelope', calibrated=False,
                                  n_members=n, n_members_present=len(values), presence_probability=1.0, status=VALUE_STATUS, source_kind=source_kind))
        for (pid, aid, ev), counts in sorted(axis_ev_count.items()):
            timings = [t for t in axis_ev_first[(pid, aid, ev)] if t is not None]
            cm, cl, ch = _quantiles([float(c) for c in counts])
            tm, tl, th = _quantiles([float(t) for t in timings]) if timings else ('unavailable',) * 3
            events.append(dict(forecast_run_id=run_id, **tag, plant_instance_id=pid, scope_id=fid_of(f'{pid}/{aid}/Axis'), event_type=ev,
                               horizon_day=h, days_after_sowing=das(h, pid), probability=r10(len(timings) / n), probability_basis=probability_basis(ev, members, (pid, aid)),
                               count_median=cm, count_low=cl, count_high=ch, timing_day_median=tm, timing_day_low=tl, timing_day_high=th, n_members=n,
                               probability_kind='member_fraction_declared_scenario_set', status=VALUE_STATUS, source_kind=source_kind, axis_id=aid,
                               truss_provenance='generated_truss' if ev == 'new_truss_generated' else 'n/a'))
        for oid, births in sorted(new_organ_birth.items()):
            med, low, high = _quantiles([float(b) for b in births])
            ev_type = f'organ_initiation:{organ_meta[oid]["cls"]}'
            pid = organ_meta[oid]['pid']
            events.append(dict(forecast_run_id=run_id, **tag, plant_instance_id=pid, scope_id=fid_of(oid),
                               event_type=ev_type, horizon_day=h, days_after_sowing=das(h, pid), probability=r10(len(births) / n),
                               probability_basis=probability_basis(ev_type, members),
                               count_median=1, count_low=1, count_high=1, timing_day_median=med, timing_day_low=low, timing_day_high=high,
                               n_members=n, probability_kind='member_fraction_declared_scenario_set', status=VALUE_STATUS, source_kind=source_kind,
                               axis_id=organ_meta[oid]['axis'], truss_provenance='generated_truss' if organ_meta[oid]['cls'] in ax.TRUSS_CLASSES else 'n/a'))
        for (pid, ev), counts in sorted(plant_counts.items()):
            timings = [t for t in plant_events[(pid, ev)] if t is not None]
            cm, cl, ch = _quantiles([float(c) for c in counts])
            tm, tl, th = _quantiles([float(t) for t in timings]) if timings else ('unavailable',) * 3
            events.append(dict(forecast_run_id=run_id, **tag, plant_instance_id=pid, scope_id=fid_of(f'{pid}/Plant'), event_type=ev,
                               horizon_day=h, days_after_sowing=das(h, pid), probability=r10(len(timings) / n), probability_basis=probability_basis(ev, members),
                               count_median=cm, count_low=cl, count_high=ch, timing_day_median=tm, timing_day_low=tl, timing_day_high=th, n_members=n,
                               probability_kind='member_fraction_declared_scenario_set', status=VALUE_STATUS, source_kind=source_kind,
                               axis_id='plant', truss_provenance='generated_truss' if ev == 'new_truss' else 'n/a'))
    return dict(forecast_mtg=mtg, forecast_intervals=intervals, forecast_events=events)


def spread_summary(intervals: list[dict], eligible: list[str], excluded: dict) -> dict:
    with_spread = [r for r in intervals if r['interval_low'] != r['interval_high']]
    return dict(n_interval_rows=len(intervals), n_rows_with_spread=len(with_spread),
                quantities_with_spread=sorted({r['quantity'] for r in with_spread}),
                ensemble_spread='none' if not eligible else ('present' if with_spread else 'eligible_but_no_numeric_spread'),
                eligible_parameters=eligible, excluded_parameters=excluded)


def scenario_comparison(intervals: list[dict], events: list[dict], last_h: int) -> dict:
    """F8 report: per-scenario per-axis leaf count / leaf area / stem length at the last horizon, stop timing, and the low-high distance."""
    out: dict[str, dict] = {}
    for r in intervals:
        if r['forecast_id'].endswith('/Axis') and int(r['horizon_day']) == last_h and r['quantity'] in ('axis_leaf_count', 'axis_leaf_area_mm2', 'axis_stem_length_mm'):
            out.setdefault(r['scenario_id'], {}).setdefault(r['axis_id'], {})[r['quantity']] = dict(base_point=r['point'], low=r['interval_low'], high=r['interval_high'])
    for e in events:
        if e['scope_id'].endswith('/Axis') and e['event_type'] == 'apex_stop' and int(e['horizon_day']) == last_h:
            out.setdefault(e['scenario_id'], {}).setdefault(e['axis_id'], {})['apex_stop'] = dict(probability=e['probability'], timing_day_median=e['timing_day_median'])
    distance = {}
    for axis in sorted({a for sc in out.values() for a in sc}):
        for q in ('axis_leaf_count', 'axis_leaf_area_mm2', 'axis_stem_length_mm'):
            lo = out.get('low', {}).get(axis, {}).get(q); hi = out.get('high', {}).get(axis, {}).get(q); ba = out.get('base', {}).get(axis, {}).get(q)
            if lo and hi and ba:
                distance[f'{axis}:{q}'] = dict(low=float(lo['base_point']), base=float(ba['base_point']), high=float(hi['base_point']),
                                               high_minus_low=r10(float(hi['base_point']) - float(lo['base_point'])),
                                               relative_to_base=r10((float(hi['base_point']) - float(lo['base_point'])) / float(ba['base_point'])) if float(ba['base_point']) else 'unavailable')
    return dict(per_scenario_per_axis_at_last_horizon=out, low_high_distance=distance,
                note='scenario range (declared set), not an uncertainty band; base is the reference; the inner envelope per scenario is the member spread')


def assumptions_table(registry: list[dict], overrides: dict, base_request: dict, scenario: dict, light: dict, pointer: dict) -> list[dict]:
    out = []
    for r in registry:
        out.append(dict(symbol=r['key'], value=json.dumps(r['value']), unit=r['unit'], source_kind=r['source_kind'], source_ref=r['source_ref'],
                        cultivar=r['cultivar'], uncertainty=r['uncertainty'], engine_key=r['engine_key'], engine_status=r['engine_status'], origin='registry'))
    for key, p in retained_from_request(base_request, overrides).items():
        out.append(dict(symbol=key, value=json.dumps(p['value']), unit=p['unit'], source_kind=p['source_kind'], source_ref=p['source_ref'],
                        cultivar=p['cultivar'], uncertainty=p['uncertainty'], engine_key=key, engine_status='v3_default_retained', origin='v3_default_retained'))
    for lamp in scenario['lamps']:
        for k in ('x_m', 'y_m', 'z_m', 'power_umol_s'):
            out.append(dict(symbol=f'LAMP:{lamp["lamp_id"]}:{k}', value=json.dumps(lamp[k]), unit=k.rsplit('_', 1)[-1], source_kind='user_assumed',
                            source_ref=f'scenario lamp {lamp["lamp_id"]} ({lamp["source_kind"]}); placement is a declared guess', cultivar='n/a',
                            uncertainty='unavailable', engine_key='lamps.csv', engine_status='environment', origin='scenario'))
    out.append(dict(symbol='REFERENCE_PPFD', value=json.dumps(light['reference_ppfd_umol_m2_s']), unit='umol_m2_s', source_kind=light['source_kind'],
                    source_ref=f'light_input.csv (measured_at={light["measured_at"] or "none"}); {light["absolute_status"]}', cultivar='n/a',
                    uncertainty='unavailable until a PPFD meter reading exists', engine_key='light_input.csv', engine_status='environment', origin='light_input'))
    out.append(dict(symbol='DLI_DERIVED', value=json.dumps(r10(light['dli_mol_m2_d'])), unit='mol_m2_d', source_kind=light['source_kind'],
                    source_ref=f'PPFD x 3600 x {light["photoperiod_h"]} h / 1e6 (derived from REFERENCE_PPFD and CLIMATE_PHOTOPERIOD_H)', cultivar='n/a',
                    uncertainty='inherits REFERENCE_PPFD', engine_key='', engine_status='layer_only', origin='derived'))
    out.append(dict(symbol='RATIO_POINTER_TARGET', value=json.dumps(pointer['target_sha256']), unit='sha256', source_kind='user_stated',
                    source_ref=f'{pointer["pointer_text"]} ({pointer["schema"]} schema, {pointer["n_rows"]} rows); role {pointer["role"]}', cultivar='n/a',
                    uncertainty='app_ratio (uncalibrated phone app, relative only)', engine_key='light_ratio_latest.txt', engine_status='environment', origin='pointer'))
    for k in ('rays', 'start_time'):
        out.append(dict(symbol=f'SCENARIO:{k}', value=json.dumps(scenario[k]), unit='count' if k == 'rays' else 'iso8601', source_kind='user_assumed',
                        source_ref='forecast scenario (numerical control, not a biological coefficient)', cultivar='n/a', uncertainty='n/a',
                        engine_key='runtime' if k == 'rays' else 'request.start_time', engine_status='environment', origin='scenario'))
    out.append(dict(symbol='SCENARIO:ensemble', value=json.dumps(scenario['ensemble']), unit='members,seed', source_kind='user_assumed',
                    source_ref='declared scenario set over eligible registry candidate_set rows; not a calibrated distribution', cultivar='n/a', uncertainty='n/a',
                    engine_key='', engine_status='layer_only', origin='scenario'))
    return out


def unresolved_table(light: dict, overrides: dict, baseline_status: str, sowing: dict, engine: str) -> list[dict]:
    u = 'unavailable'
    rows_ = [
        ('direction_change_extrapolated', 'all organs', 'mechanism (b) needs two phase-matched observations of the same organ in a common frame; none exist', u, 'forecast spec §6b'),
        ('tropism', 'all organs', 'phototropism/gravitropism not_modelled (no absolute PPFD, no light-direction description, no depth)', 'not_modelled', 'forecast spec §6c'),
        ('curvature_change', 'CompoundLeaf', 'no measured rachis theta exists; flat/bent thresholds on hold', u, 'forecast spec §5; MTG §13'),
        ('unfolding', 'CompoundLeaf', 'no recorded variable defines the stage', u, 'forecast spec §5'),
        ('organ_loss', 'all organs', 'abscission/removal unreachable: manual_events.csv and organ_label_map.csv absent (CR-5)', u, 'forecast spec §5, §13 CR-5'),
        ('side_shoot_emergence', 'plant', f'topology_scope = {TOPOLOGY_SCOPE}; whole-plant completeness unavailable', u, 'forecast spec §8'),
        ('blade_curl', 'Leaflet', 'on hold (MTG §8.3/§13): no field, no threshold', u, 'MTG §13'),
        ('rachis_bend_type', 'CompoundLeaf', 'human label only; no numeric threshold (MTG §13)', u, 'MTG §13'),
        ('fruit_set_yield', 'TrussAxis/Flower', 'outside V3 scope (flowers are slender proxies; no fruit pools)', u, 'V3 contract'),
        ('interval_coverage_accuracy', 'all', 'growth-accuracy criteria not testable until >= 3 capture dates of the same identified plants exist', 'not_testable', 'forecast spec §12 growth-accuracy'),
        ('baselines', 'persistence/recent_change', baseline_status, u, 'forecast spec §11'),
    ]
    if engine == 'v3_direct':
        rows_.append(('canopy_shading', 'scene', 'v3_direct engine: single plant pair, no P1 canopy placement; use the pipeline engine for shading', u, 'fix prompt B1'))
    for item in overrides.get('blocked_engine_extension', []):
        rows_.append((item['key'], 'engine', f'blocked_engine_extension ({item["engine_key"]}): {item["reason"]}', u, 'P3 V3 F1'))
    rows_.append(('truss_location', 'TrussAxis', f'truss geometry is a declared supplement ({ax.ZERO_EXTENT_LABEL}); A5 trusses carry no geometry; never an observed location', 'assumed', 'P3 V3 F2'))
    if sowing['sowing_source'] in ('unavailable', 'assumed_issue_minus_days_at_start'):
        rows_.append(('sowing_datetime', 'plant', f'{sowing["sowing_source"]}: {sowing.get("note", "")}', u if sowing['sowing_source'] == 'unavailable' else 'assumed', 'fix prompt B5'))
    if light['conditional']:
        rows_.append(('absolute_ppfd_dli', 'scene', f'reference PPFD is {light["source_kind"]}; every carbon-dependent value is conditional', 'conditional', 'fork Task 3.2'))
    for item in overrides['inactive_in_engine']:
        rows_.append((item['key'], 'engine', f'{item["activation"]}: {item["reason"]}', 'inactive_in_engine', 'fork Task 1-2 / fix prompt B6'))
    for item in overrides['engine_default_retained']:
        rows_.append((item['key'], 'engine', f'engine_default_retained ({item["engine_key"]}): {item["reason"]}', u, 'fork Task 1.5'))
    return [dict(quantity=q, scope=s, reason=r, status=st, spec_reference=ref) for q, s, r, st, ref in rows_]


# ---------------------------------------------------------------------------------------------------------------
# Issue
# ---------------------------------------------------------------------------------------------------------------

def issue(scenario: dict, registry_path: Path = REGISTRY_PATH, out_root: Path = ROOT / 'build', java: str = 'java',
          timeout_s: float = 1200, date: str | None = None, existing_results: list[Path] | None = None,
          engine_params: set[str] | None = None, enable_hooks: bool = False) -> Path:
    """Issue one forecast: capability check -> eligibility -> members -> engine runs (or verified results) -> aggregate -> freeze."""
    scenario = validate_scenario(scenario)
    registry_static = load_registry(registry_path)
    registry, capabilities = resolve_engine_capabilities(registry_static, engine_params, allow_hook_activation=enable_hooks)  # B6
    overrides0 = build_request_overrides(registry)
    d0 = days_at_start(overrides0)
    issue_boundary = d0 if scenario['route'] == 'seed_route' else 0
    folder = new_forecast_dir(out_root, date)
    try:
        return _issue_into(folder, scenario, registry, registry_path, overrides0, capabilities, d0, issue_boundary, java, timeout_s, existing_results)
    except Exception as exc:  # a failed issue must never look like a forecast: no manifest, explicit marker
        write_json(folder / 'FAILED.json', dict(status='failed', error=f'{type(exc).__name__}: {exc}', failed_at=now_iso(),
                                                 note='no manifest.json was written; this folder is not an issued forecast'))
        raise


def _issue_into(folder: Path, scenario: dict, registry: list[dict], registry_path: Path, overrides0: dict, capabilities: dict, d0: int,
                issue_boundary: int, java: str, timeout_s: float, existing_results: list[Path] | None) -> Path:
    from tools.microtom_forecast_eval import run_baselines
    run_id = folder.name
    engine = scenario['engine']
    write_json(folder / 'scenario.json', scenario)
    size, seed = scenario['ensemble']['size'], scenario['ensemble']['seed']
    # B3 — which declared uncertainties can act on this input state within the horizon
    controls, start_interval, days = eligibility_inputs(scenario, folder, overrides0, d0)
    elig = eligibility(registry, controls, environment_rows(overrides0, days), plain_parameters(overrides0), start_interval, days, scenario['route'])
    eligible = sorted(k for k, v in elig.items() if v['eligible'])
    excluded = {k: v['reason'] for k, v in elig.items() if not v['eligible']}
    effective_size = size if eligible else 1
    write_json(folder / 'eligibility.json', dict(controls=controls, start_interval=start_interval, days=days, report=elig, eligible=eligible,
                                                  excluded=excluded, requested_members=size, effective_members=effective_size,
                                                  rule='thermal_schedule replica of vtc_development_native.rgg::develop; a candidate is eligible iff it changes the schedule or the pair count of a compound initiated within the horizon'))
    members_meta, members = [], []
    transients: list[dict] = []  # OneDrive/Windows rename denials (WinError 5) inside engine publication; member rebuilt from scratch
    # F8: a declared scenario set applies when the engine route carries axis records (the parameter is per axis); one forecast per value
    sset = scenario_set(registry)
    axis_route = engine == 'pipeline' and bool(controls) and any(c.get('axis_id') for c in controls.values())
    base_only = scenario.get('scenario_set_mode', 'all') == 'base_only'
    if sset and axis_route and existing_results is None and not base_only:
        scenarios_run = {label: sset['values'][label] for label in SCENARIO_LABELS}
    else:
        scenarios_run = {'base': sset['values']['base'] if sset else 'n/a'}
    members_by_scenario: dict[str, list[dict]] = {}
    if existing_results is None:
        for label, sval in scenarios_run.items():
            scen_vals = {sset['parameter']: sval} if (sset and axis_route) else None
            members_by_scenario[label] = []
            for i in range(effective_size):
                values = sample_member(registry, seed, i, eligible=set(eligible))  # the SAME member draw inside every scenario
                if engine == 'pipeline':
                    sub = label if len(scenarios_run) > 1 else None
                    built = build_pipeline_member_with_retry(folder, scenario, registry, i, values, timeout_s, scen_vals, sub, transients)
                    mem = load_pipeline_member(built); members_by_scenario[label].append(mem)
                    rel = f'scenarios/{label}/members/m{i:02d}' if sub else f'members/m{i:02d}'
                    members_meta.append(dict(scenario_id=label, scenario_value=sval, member_index=i, trajectory_id=f'm{i:02d}', member_values=json.dumps(values), run_id=f'{run_id}_{label}_m{i:02d}',
                                             inputs_dir=f'{rel}/scenario', run_dir=f'{rel}/run', results_dir=f'{rel}/run/committed',
                                             results_manifest_sha256=sha(built['run_dir'] / 'collection.json'), member_source='pipeline_committed_checkpoints', engine=engine,
                                             instance_ids=json.dumps(mem['plants']), checkpoint_hashes=json.dumps([c['checkpoint_sha256'] for c in mem['checkpoints']])))
                    continue
                built = build_member_inputs(folder, scenario, registry, i, values, issue_boundary)
                results = run_member(built['inputs'], folder / 'members' / f'm{i:02d}' / 'run', java, timeout_s)
                mem = load_member_results(results, built['request']); members_by_scenario[label].append(mem)
                members_meta.append(dict(scenario_id=label, scenario_value=sval, member_index=i, trajectory_id=f'm{i:02d}', member_values=json.dumps(built['member_values']), run_id=built['run_id'],
                                         inputs_dir=f'members/m{i:02d}/inputs', run_dir=f'members/m{i:02d}/run', results_dir=f'members/m{i:02d}/run/results',
                                         results_manifest_sha256=sha(results / 'manifest.json'), member_source='native_v3_run', engine=engine,
                                         instance_ids=json.dumps(mem['plants']), checkpoint_hashes=json.dumps([])))
        members = members_by_scenario['base']
    if existing_results is not None:
        require(len(existing_results) == effective_size, f'ensemble size must equal the number of supplied result runs ({effective_size} effective)')
        for i, run_dir in enumerate(existing_results):
            run_dir = Path(run_dir)
            if engine == 'pipeline':
                scen = read_json(run_dir / 'scenario.json')
                require(NO_REUSE_MARK not in scen.get('initialization_provenance', ''),
                        f'{run_dir} is a stress-fixture run ({NO_REUSE_MARK}); it can never be a forecast member')
                coll = read_json(run_dir / 'collection.json')
                require(coll.get('status') == 'complete_conditional_trajectory' and coll['last_valid_boundary'] >= max(scenario['horizons_days']), 'pipeline result incomplete/too short')
                if i == 0 and scen['start_time'] != scenario.get('start_time'):
                    scenario = dict(scenario, start_time=scen['start_time'], start_time_status='taken_from_existing_pipeline_results (synthetic clock)')
                    write_json(folder / 'scenario.json', scenario)
                built = dict(run_dir=run_dir, scenario=scen, collection=coll, intervals=max(scenario['horizons_days']))
                members.append(load_pipeline_member(built))
                members_meta.append(dict(scenario_id='base', scenario_value=scenarios_run['base'], member_index=i, trajectory_id=f'm{i:02d}', member_values=json.dumps({}), run_id=run_dir.name,
                                         inputs_dir=str(run_dir / 'native_inputs'), run_dir=str(run_dir), results_dir=str(run_dir / 'committed'),
                                         results_manifest_sha256=sha(run_dir / 'collection.json'), member_source='existing_pipeline_committed_checkpoints', engine=engine,
                                         instance_ids=json.dumps(members[-1]['plants']), checkpoint_hashes=json.dumps([c['checkpoint_sha256'] for c in members[-1]['checkpoints']])))
                continue
            req = read_json(run_dir / 'inputs' / 'request.json')
            if i == 0 and aware(req['start_time']) != aware(scenario['start_time']):
                scenario = dict(scenario, start_time=req['start_time'], start_time_status='taken_from_existing_results_request (synthetic clock)')
                write_json(folder / 'scenario.json', scenario)
            require(aware(req['start_time']) == aware(scenario['start_time']), 'existing result runs must share one start_time')
            require(req['initialization_route'] == ('seed' if scenario['route'] == 'seed_route' else 'stage2_observation_fixture'), 'result route mismatch')
            require(req['intervals'] >= issue_boundary + max(scenario['horizons_days']), 'result run too short for the horizons')
            require(read_json(run_dir / 'execution.json')['status'] == 'verified', 'unverified result run')
            members.append(load_member_results(run_dir / 'results', req))
            diff = {k: v['value'] for k, v in req['parameters'].items() if v['value'] != members[0]['request']['parameters'][k]['value']}
            members_meta.append(dict(scenario_id='base', scenario_value=scenarios_run['base'], member_index=i, trajectory_id=f'm{i:02d}', member_values=json.dumps(dict(engine_parameters_differing_from_m00=diff)), run_id=req['run_id'],
                                     inputs_dir=str(run_dir / 'inputs'), run_dir=str(run_dir), results_dir=str(run_dir / 'results'),
                                     results_manifest_sha256=sha(run_dir / 'results' / 'manifest.json'), member_source='existing_verified_results', engine=engine,
                                     instance_ids=json.dumps(members[-1]['plants']), checkpoint_hashes=json.dumps([])))
    if existing_results is not None:
        members_by_scenario = {'base': members}
    base_request = members[0]['request']
    frame = base_request['frame']
    if engine == 'pipeline':
        scenario = dict(scenario, start_time=base_request['start_time'])
        write_json(folder / 'scenario.json', scenario)
        inputs0 = Path(members_meta[0]['inputs_dir']) if existing_results else folder / members_meta[0]['run_dir'] / 'native_inputs'
    else:
        inputs0 = Path(members_meta[0]['inputs_dir']) if existing_results else folder / members_meta[0]['inputs_dir']
    sowing = resolve_sowing(scenario, d0)
    light = load_light_input(inputs0 / 'light_input.csv', environment_values(overrides0)['light_hours']['value'])
    pointer = resolve_ratio_pointer(inputs0 / 'light_ratio_latest.txt')
    boundary_fn = lambda h: aware(scenario['start_time']) + timedelta(days=issue_boundary + h)
    das = days_after_sowing_fn(sowing, boundary_fn)
    tables = dict(forecast_mtg=[], forecast_intervals=[], forecast_events=[])
    multi = len(members_by_scenario) > 1
    for label, mems in members_by_scenario.items():
        part = aggregate(run_id, mems, scenario['horizons_days'], issue_boundary, das, frame,
                         scenario_id=label if multi else None, scenario_value=scenarios_run[label])
        for k in tables:
            tables[k].extend(part[k])
    write_table(folder / 'forecast_mtg.csv', tables['forecast_mtg'], MTG_COLUMNS)
    write_table(folder / 'forecast_intervals.csv', tables['forecast_intervals'], INTERVAL_COLUMNS)
    write_table(folder / 'forecast_events.csv', tables['forecast_events'], EVENT_COLUMNS)
    write_table(folder / 'forecast_assumptions.csv', assumptions_table(registry, overrides0, base_request, scenario, light, pointer), ASSUMPTION_COLUMNS)
    write_table(folder / 'forecast_map.csv', [], MAP_COLUMNS)
    write_table(folder / 'members.csv', members_meta, MEMBER_COLUMNS)
    carbon_cap = None
    if engine == 'pipeline':
        trajectory = [row for label, mems in members_by_scenario.items() for i, mem in enumerate(mems)
                      for row in carbon_trajectory(run_id, label, scenarios_run[label], i, mem, mem['request']['runtime']['rays'])]
        write_table(folder / 'carbon_trajectory.csv', trajectory, TRAJECTORY_COLUMNS)
        carbon_cap = cap_summary(trajectory, members_by_scenario, max(scenario['horizons_days']))
    if engine == 'pipeline':
        prepared0 =read_json((Path(members_meta[0]['run_dir']) if existing_results else folder / members_meta[0]['run_dir']) / 'prepared.json')
        snapshot = digest(dict(p1_hashes=prepared0['p1_hashes'], scenario_sha256=prepared0['scenario_sha256'], state_source=scenario['state_source']))
        snapshot_kind = f'p1_package_{scenario["state_source"]}'
        snapshot_files = sorted(prepared0['p1_hashes'])
    else:
        snapshot_files_p = [inputs0 / n for n in ('request.json', 'plants.csv', 'environment.csv', 'lamps.csv', 'light_input.csv', 'light_ratio_latest.txt')]
        snapshot_kind = 'seed_route_frozen_inputs'
        if scenario['route'] == 'stage2_fixture':
            snapshot_files_p += [inputs0 / n for n in ('initial_state.csv', 'initial_controls.json', 'mtg_vertices.csv', 'mtg_edges.csv', 'structure_state.csv', 'observation_manifest.json')]
            snapshot_kind = 'stage2_observation_fixture'
        snapshot = digest({p.name: sha(p) for p in snapshot_files_p})
        snapshot_files = [p.name for p in snapshot_files_p]
    baselines = run_baselines(folder / 'baselines', dict(frame=frame, scale_status='unavailable' if scenario['route'] == 'seed_route' else 'fixture_solved',
                                                           identity_status='model_generated' if scenario['route'] == 'seed_route' else 'fixture_confirmed',
                                                           repeated_observations=0, lights_on_clock_time=scenario.get('lights_on_clock_time'), history=None))
    write_table(folder / 'unresolved.csv', unresolved_table(light, overrides0, baselines['status'], sowing, engine), UNRESOLVED_COLUMNS)
    if engine == 'pipeline':
        executions = [read_json((Path(m['run_dir']) if existing_results else folder / m['run_dir']) / 'native' / 'execution.json') for m in members_meta]
    else:
        executions = [read_json(Path(m['run_dir']) / 'execution.json') if existing_results else read_json(folder / m['run_dir'] / 'execution.json') for m in members_meta]
    rays = base_request['runtime']['rays']
    manifest = dict(
        forecast_run_id=run_id, issued_at=now_iso(), route=scenario['route'], engine=engine, engine_scope=scenario['engine_scope'],
        state_advance_authority=('tools.tomato_fsp_pipeline (prepare_run/run; results read from committed/step_NNNN only; advance_interval never called)' if engine == 'pipeline'
                                 else 'tools.vtc_development + vtc_growth_bridge (single plant pair, no canopy)'),
        issue_class=scenario['issue_class'], scenario_id=scenario['scenario_id'], scenario=scenario,
        input_snapshot_kind=snapshot_kind, input_snapshot_sha256=snapshot, input_snapshot_files=snapshot_files,
        code_sha256=code_digest([p for p in CODE_FILES if p.exists()]), code_files=[p.name for p in CODE_FILES if p.exists()], parameter_sha256=registry_hash(registry_path),
        registry_path=str(registry_path.relative_to(ROOT)) if registry_path.is_relative_to(ROOT) else str(registry_path),
        engine_capabilities=capabilities,
        ensemble=dict(requested_size=size, size=effective_size, seed=seed, sampled_rows=eligible, eligible_parameters=eligible, excluded_parameters=excluded,
                      spread_kind='declared_scenario_set', calibrated_probability=False,
                      members=[dict(index=m['member_index'], values=json.loads(m['member_values'])) for m in members_meta]),
        ensemble_spread=spread_summary([r for r in tables['forecast_intervals'] if r['scenario_id'] == 'base'], eligible, excluded),
        engine_runtime=dict(mode=base_request['mode'], groimp_version=sorted({e.get('groimp_version', 'unavailable') for e in executions}),
                            java_version=sorted({e.get('java_version', '').strip().splitlines()[0] if e.get('java_version') else 'unavailable' for e in executions}),
                            rays=rays, quantitative_light_sampling=rays >= QUANTITATIVE_RAYS and scenario['issue_class'] == 'quantitative',
                            member_source=members_meta[0]['member_source']),
        instance_ids=sorted({pid for m in members for pid in m['plants']}),
        scenario_set=(dict(sset, applied=list(scenarios_run), what_scenarios_vary=f'{sset["parameter"]} on every S axis (M and root stay 1.0)',
                           what_members_vary=f'registry candidate_set rows that can act on the input state: {eligible}', pooled=False,
                           status='three_forecasts_side_by_side' if multi else ('single_scenario_base_only: ' + ('existing results carry their own axis parameters' if existing_results else
                                                                                                                 'requested by scenario_set_mode=base_only (diagnostic of the reference scenario)' if base_only else
                                                                                                                 'no declared axes on this route, parameter not applicable')),
                           comparison=scenario_comparison(tables['forecast_intervals'], tables['forecast_events'], scenario['horizons_days'][-1]) if multi else None)
                      if sset else None),
        axis_mode='fsp-axes-1' if all(m['request'].get('axis_contract') for m in members) else 'legacy_single_axis (axis_id label M)',
        axes=[dict(plant_instance_id=a['plant_instance_id'], axis_id=a['axis_id'], order=a['order'], stage=a.get('axis_state', a.get('stage')),
                   axil_parent_rank=a['parent_rank'], producer_segment_index=a.get('segment_index', 'unavailable'),
                   truss_rule_status=members[0].get('axis_inventory', {}).get(1, {}).get((a['plant_instance_id'], a['axis_id']), {}).get('truss_rule_status', a.get('truss_rule_status', 'unavailable')),
                   axis_segment_index_at_boundary_1=members[0].get('axis_inventory', {}).get(1, {}).get((a['plant_instance_id'], a['axis_id']), {}).get('segment_index', 'unavailable'),
                   registry_truss_history=a.get('truss_history', []), apply_determinate_stop=a['parameters']['APPLY_DETERMINATE_STOP']['value'],
                   substem_sink_scale=a['parameters']['SUBSTEM_SINK_SCALE']['value']) for a in base_request.get('axes', [])],
        axis_columns=dict(axis_id='logical M/S<k> (fsp-axes-1) or the legacy label M', axil_parent_rank='side attachment rank on M (engine state parent_rank; 0 for M); NOT the truss above-rank',
                          organ_segment_index='immutable organ segment (engine state/organ_inventory segment_index)', axis_segment_index='current axis counter at the boundary (committed axis_inventory.csv segment_index = 1 + trusses)',
                          truss_provenance='registry_truss (imported from P1/registry at the issue boundary) vs generated_truss (model event)'),
        truss_geometry_status=ax.ZERO_EXTENT_LABEL if base_request.get('axis_contract') else 'n/a',
        axis_sums_check=[m.get('axis_sums') for m in members] if base_request.get('axis_contract') else [],
        transient_retries=transients,
        rerun_trigger=dict(claimed=scenario['rerun_trigger'], meaning=RERUN_TRIGGERS[scenario['rerun_trigger']],
                           conditions={k: v for k, v in RERUN_TRIGGERS.items() if k != 'initial_issue'},
                           rule='G3 (user decision 2026-10-13): re-issue on the main fixture when any condition becomes true; base stays the reference; low_high_distance is reported beside the headline'),
        hidden_state_passed_through=overrides0.get('hidden_state', {}),
        carbon_cap=carbon_cap,
        blocked_engine_extension=overrides0.get('blocked_engine_extension', []), headline=HEADLINE,
        real_plant_verified=False, real_plant_note='engine has only been run from a P1 package through the fixture_truth route; real-photo scale unclosed (0 of 117 gaps); format-verified only',
        checkpoints=[dict(member=m['member_index'], hashes=json.loads(m['checkpoint_hashes'])) for m in members_meta] if engine == 'pipeline' else [],
        horizons_days=scenario['horizons_days'], issue_boundary_day=issue_boundary, days_at_start=d0,
        sowing_datetime=sowing['sowing_datetime'], sowing_source=sowing['sowing_source'], sowing_per_plant=sowing['per_plant'],
        run_status=RUN_STATUS, evaluation_status=EVALUATION_STATUS, conditional_simulation=True, biological_validation='not_established',
        topology_scope=TOPOLOGY_SCOPE, completeness='unavailable', closure='native_conditional_v3_track_b', cr3_compound_envelope_in_light=False,
        direction_mechanisms=dict(fixed_generation='active_user_assumed', extrapolated_observed_change='unavailable', tropism='not_modelled'),
        inactive_in_engine=overrides0['inactive_in_engine'], activated_by_capability_check=overrides0['activated_by_capability_check'],
        engine_default_retained=overrides0['engine_default_retained'],
        v3_default_retained=sorted(retained_from_request(base_request, overrides0)), candidates_registered=[c['key'] for c in overrides0['candidates']],
        light=dict(reference_ppfd_umol_m2_s=light['reference_ppfd_umol_m2_s'], source_kind=light['source_kind'], dli_mol_m2_d=r10(light['dli_mol_m2_d']),
                   conditional=light['conditional'], ratio_role=pointer['role'], ratio_target=pointer['pointer_text'], ratio_target_sha256=pointer['target_sha256'],
                   lamp_placement_status='declared_guess'),
        photoperiod_phase=dict(lights_on_clock_time=scenario.get('lights_on_clock_time'),
                               time_since_lights_on_min_at_boundaries=time_since_lights_on_min(scenario['start_time'], scenario.get('lights_on_clock_time'))),
        declared_assumptions=scenario['declared_assumptions'], numeric_rounding=f'{SIGNIFICANT_DIGITS} significant digits',
        growth_accuracy_status='not_testable_until_3_capture_dates_of_the_same_plants', baselines=baselines,
        value_source_kind_note='forecast values are conditional simulation outputs under user_assumed and literature_other_cultivar parameters (forecast_assumptions.csv); none is a measurement')
    finalize(folder, manifest)
    return folder


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest='op', required=True)
    p = sub.add_parser('scenario', help='write a scenario JSON')
    p.add_argument('--route', choices=ROUTES, default=None); p.add_argument('--engine', choices=ENGINES, default=None)
    p.add_argument('--out', type=Path, required=True)
    p.add_argument('--horizons', type=int, nargs='+', default=[1, 3, 7, 14]); p.add_argument('--members', type=int, default=4)
    p.add_argument('--seed', type=int, default=11); p.add_argument('--rays', type=int, default=QUANTITATIVE_RAYS)
    p.add_argument('--diagnostic', action='store_true', help='required when --rays is below 4.8M; marks issue_class = diagnostic')
    p.add_argument('--source-run', default=None, help='stage2_fixture: verified V3 run dir relative to the project root')
    p.add_argument('--fixture-day', type=int, default=30); p.add_argument('--lights-on', default=None, help='HH:MM, user_stated; omitted = unavailable')
    p.add_argument('--p1', default=None, help='pipeline engine: P1 package dir relative to the project root'); p.add_argument('--state-source', choices=['main', 'fixture_truth'], default='fixture_truth')
    p.add_argument('--sowing-datetime', default=None, help='aware ISO8601; recorded as user_stated')
    p.add_argument('--rerun-trigger', choices=sorted(RERUN_TRIGGERS), default='initial_issue', help='G3: which re-run condition this issue claims')
    p.add_argument('--base-only', action='store_true', help='F10: issue the reference (base) scenario alone instead of the declared scenario set')
    p = sub.add_parser('issue', help='issue one forecast (write-once folder under build/)')
    p.add_argument('--rerun-trigger', choices=sorted(RERUN_TRIGGERS), default=None, help='G3: override the scenario rerun_trigger for this issue')
    p.add_argument('--base-only', action='store_true', help='F10: issue the reference (base) scenario alone')
    p.add_argument('--scenario', type=Path, required=True); p.add_argument('--registry', type=Path, default=REGISTRY_PATH)
    p.add_argument('--out-root', type=Path, default=ROOT / 'build'); p.add_argument('--java', default='java'); p.add_argument('--timeout', type=float, default=1200)
    p.add_argument('--date', default=None, help='YYYYMMDD for the folder name (default: today)')
    p.add_argument('--rays', type=int, default=None, help='override the scenario ray count (below 4.8M needs --diagnostic)'); p.add_argument('--diagnostic', action='store_true')
    p.add_argument('--from-results', type=Path, nargs='*', default=None, help='consume these verified V3 run dirs as members instead of running')
    p.add_argument('--enable-hooks', action='store_true', help='B6: allow H1-H3 activation when V3 exposes the hook and its contract is met (owner RESULT required)')
    a = parser.parse_args(argv)
    try:
        if a.op == 'scenario':
            route = a.route or ('p1_pipeline' if (a.engine == 'pipeline' or a.p1) else 'seed_route')
            if a.engine:
                require(ROUTE_ENGINE[route] == a.engine, f'--engine {a.engine} does not fit route {route}')
            s = default_scenario(route, a.horizons, a.members, a.seed, a.rays, source_run=a.source_run, fixture_day=a.fixture_day,
                                 p1_package=a.p1, state_source=a.state_source, sowing_datetime=a.sowing_datetime, diagnostic=a.diagnostic,
                                 rerun_trigger=a.rerun_trigger, base_only=a.base_only)
            s['lights_on_clock_time'] = a.lights_on
            validate_scenario(s); write_json(a.out, s); print(json.dumps(dict(scenario=str(a.out), engine=s['engine'], issue_class=s['issue_class']), indent=2))
        else:
            s = read_json(a.scenario)
            if a.rays is not None:
                s['rays'] = a.rays
                s['issue_class'] = 'diagnostic' if (a.diagnostic or a.rays < QUANTITATIVE_RAYS) else 'quantitative'
            elif a.diagnostic:
                s['issue_class'] = 'diagnostic'
            if a.rerun_trigger is not None:
                s['rerun_trigger'] = a.rerun_trigger
            if a.base_only:
                s['scenario_set_mode'] = 'base_only'
            folder = issue(s, a.registry, a.out_root, a.java, a.timeout, a.date, a.from_results, enable_hooks=a.enable_hooks)
            print(json.dumps(dict(forecast=str(folder), manifest=str(folder / 'manifest.json')), indent=2))
    except (ContractError, ValueError, KeyError, OSError, TimeoutError) as exc:
        parser.exit(1, f'{type(exc).__name__}: {exc}\n')


if __name__ == '__main__':
    main()
